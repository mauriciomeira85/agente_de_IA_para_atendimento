# ==============================================================================
# ARQUIVO: agente/nos.py
# ------------------------------------------------------------------------------
# INTRODUÇÃO
#
# Este arquivo contém os "nós" do grafo do LangGraph — cada função aqui é
# uma etapa do processamento de UMA mensagem recebida de um cliente. O
# LangGraph vai chamando essas funções em sequência (ver agente/grafo.py),
# sempre passando adiante o mesmo "estado" (EstadoConversa), que cada nó lê
# e atualiza.
#
# A ordem de execução é:
#   1) interpretar_midia        -> transforma áudio/imagem/documento em texto
#   2) executar_turno_do_agente -> chama a DeepSeek com ferramentas reais e
#                                   deixa O MODELO decidir se/quando usá-las
#
# Este mecanismo de "turno/passo" é o mesmo reaproveitado, sem alteração de
# lógica, nos três projetos da linhagem (ver Informacoes/Arquitetura.md,
# seção 3) — o modelo recebe as ações disponíveis como FERRAMENTAS de
# verdade (ver agente/ferramentas.py) e decide sozinho, no meio da própria
# resposta, se vale a pena chamar alguma. O mecanismo do loop em si mora em
# agente/loop_de_ferramentas.py (reaproveitado também pela varredura
# periódica de reengajamento, ver agente/monitoramento.py) — este arquivo
# só monta o prompt e as ferramentas específicas de UMA conversa antes de
# chamá-lo.
# ==============================================================================

import base64

import structlog

from app.agente.estados import EstadoConversa
from app.agente.ferramentas import (
    NUMERO_MAXIMO_DE_MENSAGENS_LIVRES_ANTES_DO_ENCAMINHAMENTO,
    ContextoDoTurno,
    RodadaDoTurno,
    forcar_encaminhamento_por_limite_de_mensagens,
    montar_ferramentas_de_reengajamento,
    montar_ferramentas_do_turno,
)
from app.agente.guardrails import (
    MENSAGEM_DE_FALLBACK_PARA_RESPOSTA_BLOQUEADA,
    detectar_resposta_suspeita,
)
from app.agente.loop_de_ferramentas import rodar_loop_de_ferramentas
from app.agente.prompts import montar_mensagens_para_o_modelo
from app.armazenamento_midia import salvar_midia
from app.integracoes_externas.deepseek import descrever_imagem
from app.integracoes_externas.together_transcricao import transcrever_audio
from app.integracoes_externas.video import extrair_frame_de_video
from app.integracoes_externas.whatsapp import baixar_bytes_da_midia, obter_url_de_download_da_midia
from app.midia.leitor_pdf import extrair_texto_de_pdf, renderizar_primeira_pagina_como_imagem

logger = structlog.get_logger(__name__)


async def interpretar_midia(estado: EstadoConversa) -> dict:
    """
    Nó 1: se a mensagem recebida não for texto puro, baixa o arquivo de
    verdade da WhatsApp Cloud API (a Meta só manda um ID no webhook, nunca
    o conteúdo — ver integracoes_externas/whatsapp.py) e usa o modelo
    apropriado pra transformar em texto:
      - imagem (foto/sticker) -> DeepSeek (deepseek-flash, ver
        integracoes_externas/deepseek.py:descrever_imagem) — além da
        descrição em texto (guardada pro histórico), a imagem em si é
        anexada à mensagem do turno ATUAL (ver
        imagem_para_anexar_no_turno em agente/estados.py)
      - vídeo (é assim que a Meta entrega um GIF do seletor nativo do
        WhatsApp) -> extrai um frame com ffmpeg (ver
        integracoes_externas/video.py) e segue o mesmo caminho da imagem
      - PDF -> extração de texto local primeiro (sem IA, ver
        app/midia/leitor_pdf.py:extrair_texto_de_pdf); só quando não
        sobra texto nenhum (provavelmente escaneado) é que a primeira
        página é renderizada como imagem e segue o mesmo caminho da
        imagem/vídeo acima — um cliente pode mandar uma nota fiscal, um
        comprovante ou um print em PDF durante o atendimento
      - áudio -> Whisper Large v3 via Together AI (ver
        integracoes_externas/together_transcricao.py)
      - .docx -> extração de texto direta (sem IA, ver
        _extrair_texto_de_docx abaixo)
      - .xlsx/.xlsm -> extração de texto direta (sem IA, ver
        _extrair_texto_de_xlsx abaixo) — .xls antigo (binário) não é
        suportado
      - .txt -> decodificado direto como texto

    Se a mensagem já for texto, este nó apenas repassa o texto adiante sem
    gastar uma chamada extra de IA.
    """
    if estado["tipo_mensagem_recebida"] == "texto":
        return {"mensagem_recebida_em_texto": estado["conteudo_bruto_recebido"]}

    media_id = estado.get("id_midia_whatsapp")
    if not media_id:
        logger.warning("midia_sem_media_id", tipo=estado["tipo_mensagem_recebida"])
        return {
            "mensagem_recebida_em_texto": f"[O cliente enviou um(a) {estado['tipo_mensagem_recebida']}, mas não foi possível recuperar o arquivo. Avise que não conseguiu abrir e peça para reenviar ou descrever em texto.]"
        }

    integracao = estado["integracao_whatsapp"]
    try:
        url_de_download = await obter_url_de_download_da_midia(media_id, integracao.token_de_acesso)
        conteudo_do_arquivo = await baixar_bytes_da_midia(url_de_download, integracao.token_de_acesso)
    except ValueError:
        logger.error("falha_ao_baixar_midia_recebida", media_id=media_id)
        return {
            "mensagem_recebida_em_texto": "[Não foi possível baixar o arquivo que o cliente enviou. Avise isso e peça para reenviar.]"
        }

    # Guarda uma cópia própria em disco (ver armazenamento_midia.py) — a
    # URL de download que acabamos de usar expira em minutos, então sem
    # isso a aba Conversas nunca conseguiria reexibir a mídia depois.
    salvar_midia(media_id, conteudo_do_arquivo)

    mime_type = estado.get("mime_type_da_midia") or ""
    nome_do_arquivo = estado.get("nome_do_arquivo_da_midia") or ""

    if estado["tipo_mensagem_recebida"] == "imagem":
        eh_video = mime_type.startswith("video/")

        if eh_video:
            frame = extrair_frame_de_video(conteudo_do_arquivo)
            if not frame:
                return {
                    "mensagem_recebida_em_texto": "[O cliente enviou um vídeo, mas não foi possível extrair um frame para interpretar agora. Avise isso e peça para descrever em texto ou mandar como imagem/áudio, se possível.]"
                }
            return await _interpretar_imagem_e_anexar(
                frame,
                "image/jpeg",
                rotulo_de_origem="um vídeo curto (ex.: um GIF) — mostrando um frame extraído dele",
                texto_de_falha="[O cliente enviou um vídeo, mas não foi possível interpretar o conteúdo agora. Avise isso e continue a conversa normalmente.]",
            )

        return await _interpretar_imagem_e_anexar(
            conteudo_do_arquivo,
            mime_type or "image/jpeg",
            rotulo_de_origem="uma imagem",
            texto_de_falha="[O cliente enviou uma imagem, mas não foi possível interpretar o conteúdo agora. Avise isso e continue a conversa normalmente.]",
        )

    if estado["tipo_mensagem_recebida"] == "audio":
        transcricao = await transcrever_audio(conteudo_do_arquivo, nome_do_arquivo or "audio.ogg", mime_type or "audio/ogg")
        if not transcricao:
            return {
                "mensagem_recebida_em_texto": "[O cliente enviou um áudio, mas não foi possível transcrever agora. Avise isso e peça para escrever a mensagem em texto.]"
            }
        return {"mensagem_recebida_em_texto": transcricao}

    # Documento — PDF, Word (.docx) ou texto simples (.txt), conforme a
    # extensão do nome do arquivo que a Meta manda.
    extensao = nome_do_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_do_arquivo else ""

    if extensao == "pdf" or mime_type == "application/pdf":
        texto_extraido_do_pdf = extrair_texto_de_pdf(conteudo_do_arquivo)
        if texto_extraido_do_pdf:
            return {"mensagem_recebida_em_texto": f"[O cliente enviou um documento PDF. Conteúdo: {texto_extraido_do_pdf}]"}

        texto_de_falha_do_pdf = "[O cliente enviou um PDF, mas não foi possível ler o conteúdo agora. Avise isso e peça um resumo em texto, se for importante para a conversa.]"
        pagina_como_imagem = renderizar_primeira_pagina_como_imagem(conteudo_do_arquivo)
        if not pagina_como_imagem:
            return {"mensagem_recebida_em_texto": texto_de_falha_do_pdf}

        return await _interpretar_imagem_e_anexar(
            pagina_como_imagem,
            "image/png",
            rotulo_de_origem="um documento PDF (provavelmente escaneado) — mostrando a primeira página",
            texto_de_falha=texto_de_falha_do_pdf,
        )

    if extensao == "docx":
        texto_extraido = _extrair_texto_de_docx(conteudo_do_arquivo)
        if not texto_extraido:
            return {
                "mensagem_recebida_em_texto": "[O cliente enviou um documento Word vazio ou ilegível. Avise isso normalmente.]"
            }
        return {"mensagem_recebida_em_texto": f"[O cliente enviou um documento Word. Conteúdo: {texto_extraido}]"}

    if extensao in ("xlsx", "xlsm") or mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        texto_extraido = _extrair_texto_de_xlsx(conteudo_do_arquivo)
        if not texto_extraido:
            return {
                "mensagem_recebida_em_texto": "[O cliente enviou uma planilha Excel vazia, ilegível ou num formato antigo (.xls) não suportado — peça pra reenviar em .xlsx, se possível.]"
            }
        return {"mensagem_recebida_em_texto": f"[O cliente enviou uma planilha Excel. Conteúdo: {texto_extraido}]"}

    if extensao == "txt":
        try:
            texto_extraido = conteudo_do_arquivo.decode("utf-8", errors="ignore").strip()
        except Exception:  # noqa: BLE001 — arquivo de texto malformado não deve travar a conversa
            texto_extraido = ""
        if not texto_extraido:
            return {"mensagem_recebida_em_texto": "[O cliente enviou um arquivo de texto vazio ou ilegível.]"}
        return {"mensagem_recebida_em_texto": f"[O cliente enviou um arquivo de texto. Conteúdo: {texto_extraido}]"}

    return {
        "mensagem_recebida_em_texto": f"[O cliente enviou um arquivo ({nome_do_arquivo or 'sem nome'}) de um tipo não suportado. Avise isso e peça para descrever o conteúdo em texto.]"
    }


async def _interpretar_imagem_e_anexar(conteudo_da_imagem: bytes, mime_type: str, rotulo_de_origem: str, texto_de_falha: str) -> dict:
    """
    Descreve uma imagem (foto, sticker, frame de vídeo ou página de PDF
    renderizada) via DeepSeek e devolve tanto a descrição em texto — que
    vira o histórico reenviado em turnos FUTUROS — quanto a própria
    imagem, pronta pra ser anexada à mensagem do turno ATUAL.
    """
    descricao = await descrever_imagem(conteudo_da_imagem, mime_type)
    if not descricao:
        return {"mensagem_recebida_em_texto": texto_de_falha}

    imagem_em_base64 = base64.b64encode(conteudo_da_imagem).decode("ascii")
    return {
        "mensagem_recebida_em_texto": f"[O cliente enviou {rotulo_de_origem}. Descrição automática: {descricao}]",
        "imagem_para_anexar_no_turno": f"data:{mime_type};base64,{imagem_em_base64}",
    }


def _extrair_texto_de_docx(conteudo_do_arquivo: bytes) -> str:
    """Extrai o texto puro de um documento Word (.docx) recebido em conversa — não usa IA nenhuma."""
    import io

    from docx import Document

    documento = Document(io.BytesIO(conteudo_do_arquivo))
    paragrafos = [paragrafo.text.strip() for paragrafo in documento.paragraphs if paragrafo.text.strip()]
    return "\n".join(paragrafos)


def _extrair_texto_de_xlsx(conteudo_do_arquivo: bytes) -> str:
    """Extrai o conteúdo de uma planilha Excel (.xlsx/.xlsm) recebida em conversa como texto simples — não usa IA nenhuma."""
    import io

    from openpyxl import load_workbook

    try:
        pasta_de_trabalho = load_workbook(io.BytesIO(conteudo_do_arquivo), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 — arquivo corrompido, .xls antigo ou outro formato não deve travar a conversa
        return ""

    linhas_de_texto: list[str] = []
    for nome_da_planilha in pasta_de_trabalho.sheetnames:
        for linha in pasta_de_trabalho[nome_da_planilha].iter_rows(values_only=True):
            celulas = [str(valor).strip() for valor in linha if valor is not None and str(valor).strip()]
            if celulas:
                linhas_de_texto.append(" | ".join(celulas))

    return "\n".join(linhas_de_texto)


async def executar_turno_do_agente(estado: EstadoConversa) -> dict:
    """
    Nó 2: monta o prompt e as ferramentas disponíveis, e roda o loop de
    turno/passo (ver agente/loop_de_ferramentas.py) — o mesmo mecanismo
    genérico reaproveitado pela varredura periódica de reengajamento (ver
    agente/monitoramento.py), aqui aplicado a UMA conversa.

    Lê `estado["sessao"]` (a Session do SQLAlchemy) — diferente dos outros
    dois projetos da linhagem, as ferramentas deste domínio
    (consultar_base_de_conhecimento, encaminhar_para_setor) precisam fazer
    consultas ao banco DENTRO do próprio turno (busca vetorial, validação
    de setor — ver agente/ferramentas.py:ContextoDoTurno). Vem do estado
    (não de um parâmetro extra da função) porque o LangGraph só passa o
    `estado` para cada nó, sem argumentos adicionais.
    """
    perfil = estado["perfil_da_empresa"]
    rodada = RodadaDoTurno()
    contexto = ContextoDoTurno(
        perfil=perfil,
        rodada=rodada,
        atendimento=estado["atendimento"],
        empresa=estado["empresa"],
        configuracao=estado["configuracao_agente"],
        integracao_whatsapp=estado["integracao_whatsapp"],
        sessao=estado["sessao"],
    )

    # Se o atendimento já tinha sido encaminhado ANTES desta mensagem, a
    # ação já aconteceu de verdade numa mensagem passada — não faz sentido
    # oferecer a mesma ferramenta de novo a cada "obrigado"/"até" que o
    # cliente manda depois. Mas se ele trouxer algo NOVO E RELEVANTE, pode
    # ser reencaminhado — hoje UMA ÚNICA VEZ por conversa (ver
    # Atendimento.encaminhamento_definitivo em agente/orquestrador.py, que
    # bloqueia qualquer turno futuro antes mesmo de chegar aqui).
    atendimento_ja_encaminhado = estado.get("status_atual_atendimento") == "encaminhado"
    numero_de_mensagens_livres_ja_enviadas = estado.get("mensagens_livres_ja_enviadas", 0)

    mensagens = montar_mensagens_para_o_modelo(
        perfil=perfil,
        historico=estado["historico_mensagens"],
        mensagem_atual=estado["mensagem_recebida_em_texto"],
        atendimento_ja_encaminhado=atendimento_ja_encaminhado,
        imagem_anexada=estado.get("imagem_para_anexar_no_turno"),
        mensagens_livres_ja_enviadas=numero_de_mensagens_livres_ja_enviadas,
    )

    if atendimento_ja_encaminhado:
        ferramentas = montar_ferramentas_de_reengajamento(contexto)
        texto_final = await rodar_loop_de_ferramentas(
            mensagens,
            ferramentas,
            id_para_log=estado.get("id_atendimento"),
            ferramentas_terminais=frozenset({"nao_responder"}),
            ferramentas_de_fechamento=frozenset({"encaminhar_para_setor"}),
        )
    else:
        ferramentas = montar_ferramentas_do_turno(contexto)
        texto_final = await rodar_loop_de_ferramentas(
            mensagens,
            ferramentas,
            id_para_log=estado.get("id_atendimento"),
            ferramentas_terminais=frozenset({"nao_responder"}),
            ferramentas_de_fechamento=frozenset({"marcar_como_resolvido", "encaminhar_para_setor"}),
        )

        # Rede de segurança do teto de custo por atendimento: esta era a
        # ÚLTIMA mensagem livre permitida e, mesmo assim, o modelo
        # respondeu em texto normal SEM chamar nenhuma ferramenta de ação
        # — uma instrução de prompt nunca é 100% garantida. Força o
        # encaminhamento agora, em código, para que o teto nunca seja
        # ultrapassado de verdade. texto_final is not None exclui o caso
        # de o modelo ter chamado "nao_responder" (silêncio proposital,
        # nada a forçar).
        ja_no_teto = (
            numero_de_mensagens_livres_ja_enviadas
            >= NUMERO_MAXIMO_DE_MENSAGENS_LIVRES_ANTES_DO_ENCAMINHAMENTO - 1
        )
        if ja_no_teto and texto_final is not None and rodada.desfecho_acionado is None:
            await forcar_encaminhamento_por_limite_de_mensagens(contexto)
            texto_final = (
                f"{texto_final}\n\nVou te colocar em contato direto com o nosso time agora, combinado?"
            )

    logger.info(
        "turno_do_agente_concluido",
        id_atendimento=estado.get("id_atendimento"),
        desfecho=rodada.desfecho_acionado,
    )

    if texto_final is None:
        # SILÊNCIO PROPOSITAL: o modelo chamou "nao_responder", em
        # qualquer ponto da conversa — diferente do teto de passos do loop
        # (que devolve "", tratado abaixo). agente/orquestrador.py sabe
        # que None significa não enviar nada ao cliente neste turno.
        resposta_para_envio = None
    else:
        resposta_para_envio = texto_final or "Só um momento, já te retorno com mais detalhes!"

        # Filtro de saída (guardrail): revisa a resposta que o modelo
        # escreveu ANTES de enviar ao cliente, procurando sinais de que uma
        # tentativa de manipulação (prompt injection) funcionou.
        padrao_suspeito = detectar_resposta_suspeita(resposta_para_envio)
        if padrao_suspeito:
            logger.warning(
                "resposta_bloqueada_por_guardrail",
                id_atendimento=estado.get("id_atendimento"),
                padrao_detectado=padrao_suspeito,
                resposta_original=resposta_para_envio,
            )
            resposta_para_envio = MENSAGEM_DE_FALLBACK_PARA_RESPOSTA_BLOQUEADA

    return {
        "resposta_para_envio": resposta_para_envio,
        "desfecho_acionado": rodada.desfecho_acionado,
    }


# ==============================================================================
# RESUMO
# ------------------------------------------------------------------------------
# Este arquivo implementa os dois nós do grafo do agente: interpretar_midia
# (baixa a mídia de verdade da WhatsApp Cloud API e traduz áudio/imagem/
# vídeo/PDF/Word/Excel/txt para texto) e executar_turno_do_agente (monta um
# ContextoDoTurno com o atendimento/empresa/credenciais/sessão reais e roda
# o loop de turno com ferramentas reais, deixando o modelo decidir sozinho
# se/quando/qual ação chamar — inclusive consultando a Base de Conhecimento
# e encaminhando de verdade para um setor). Diferente dos outros dois
# projetos, este nó recebe a `sessao` do banco diretamente, porque as
# ferramentas deste domínio precisam consultar o banco dentro do próprio
# turno (busca vetorial, validação de setor). "nao_responder" está
# disponível em QUALQUER turno. Antes do encaminhamento, o nó também
# aplica a trava de custo por atendimento (mesmo padrão dos outros dois
# projetos). Antes de devolver qualquer resposta em texto, o guardrail de
# saída revisa o texto em busca de sinais de manipulação.
# ==============================================================================
